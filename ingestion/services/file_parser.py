import csv
import io
from datetime import date, datetime


class FileError(Exception):
    """A problem with the whole file (wrong type, unreadable, empty), shown to the user as is."""


class FileParser:
    """Reads a CSV or .xlsx payroll file into header names and rows of text.

    Every cell comes back as text (or a date for an Excel date cell), so nothing is guessed:
    an account number like 0123456789 keeps its leading zero.
    """

    @staticmethod
    def parse(file) -> tuple[list[str], list[list]]:
        name = (file.name or "").lower()
        if name.endswith(".csv"):
            return FileParser._read_csv(file.read())
        if name.endswith(".xlsx"):
            return FileParser._read_xlsx(file)
        if name.endswith(".xls"):
            raise FileError("Old .xls files are not supported. Save the file as .xlsx or .csv and upload it again.")
        raise FileError("Upload a .csv or .xlsx file.")

    @staticmethod
    def _read_csv(raw: bytes):
        # Excel's "CSV UTF-8" adds a BOM; its plain "CSV" uses the Windows-1252 code page.
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise FileError("The CSV file could not be read. Save it as CSV UTF-8 and upload it again.")
        rows = list(csv.reader(io.StringIO(text)))
        return FileParser._split(rows)

    @staticmethod
    def _read_xlsx(file):
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(file, read_only=True, data_only=True)
        except Exception:
            raise FileError("The .xlsx file could not be opened. Check that it is a real Excel file.")
        try:
            sheet = workbook.worksheets[0]
            rows = [[FileParser._cell_text(value) for value in row] for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
        return FileParser._split(rows)

    @staticmethod
    def _cell_text(value):
        if value is None:
            return ""
        if isinstance(value, (datetime, date)):
            return value  # A real Excel date; the standardiser turns it into a date.
        if isinstance(value, float) and value.is_integer():
            return str(int(value))  # 250000.0 -> "250000"
        return str(value)

    @staticmethod
    def _split(rows):
        if not rows or not any(str(cell).strip() for cell in rows[0]):
            raise FileError("The file is empty. The first row must hold the column names.")
        headers = [str(cell) for cell in rows[0]]
        return headers, rows[1:]
