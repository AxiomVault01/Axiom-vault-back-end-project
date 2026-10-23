from datetime import timedelta
from io import BytesIO

from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from accounts.models import AuditLog
from fraud.models import Case, CaseComment, CaseEvent, CaseEvidence
from fraud.services.report import ReportService
from fraud.test_cases import client_for, make_alert, member, org, outsider, owner  # noqa: F401 (fixtures)
from ingestion.models import StoredFile

CASES = "/api/v1/fraud/cases/"


def closed_case(organisation, number=1, **extra):
    values = {"title": f"Case {number}", "priority": "high", "status": "closed", "closed_at": timezone.now() - timedelta(days=3)}
    values.update(extra)
    return Case.objects.create(organisation=organisation, number=number, **values)


def set_status(user, reference, status):
    return client_for(user).post(f"{CASES}{reference}/status/", {"status": status}, format="json")


def test_archive_a_closed_case_keeps_closed_at(owner, org):
    case = closed_case(org)
    closed_at = case.closed_at

    response = set_status(owner, "CASE-0001", "archived")

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "archived"
    assert body["sla_days_remaining"] is None
    case.refresh_from_db()
    assert case.closed_at == closed_at


def test_only_a_closed_case_can_be_archived(owner, org):
    Case.objects.create(organisation=org, number=1, title="Open", priority="high", status="investigating")

    response = set_status(owner, "CASE-0001", "archived")

    assert response.status_code == 400
    assert response.json() == {"status": ["Only a closed case can be archived. Close it first."]}


def test_restore_an_archived_case(owner, org):
    case = closed_case(org, status="archived")
    closed_at = case.closed_at

    restored = set_status(owner, "CASE-0001", "closed")
    assert restored.status_code == 200
    case.refresh_from_db()
    assert (case.status, case.closed_at) == ("closed", closed_at)

    set_status(owner, "CASE-0001", "archived")
    reopened = set_status(owner, "CASE-0001", "investigating").json()
    assert (reopened["status"], reopened["closed_at"]) == ("investigating", None)
    assert set_status(owner, "CASE-0001", "escalated").status_code == 400


def test_an_archived_case_cannot_be_escalated(owner, member, org):
    closed_case(org, status="archived")

    response = client_for(owner).post(
        f"{CASES}CASE-0001/escalate/",
        {"reason": "legal_risk", "escalate_to_id": str(member.id), "priority": "high"},
        format="json",
    )

    assert response.status_code == 400
    assert response.json() == {"status": ["An archived case cannot be escalated. Restore it first."]}


def test_archived_cases_have_their_own_list(owner, org):
    closed_case(org, 1)
    closed_case(org, 2, status="archived")
    Case.objects.create(organisation=org, number=3, title="Open", priority="low")
    client = client_for(owner)

    main = [c["reference"] for c in client.get(CASES).json()["results"]]
    archived = [c["reference"] for c in client.get(CASES, {"status": "archived"}).json()["results"]]

    assert sorted(main) == ["CASE-0001", "CASE-0003"]
    assert archived == ["CASE-0002"]
    assert client.get(f"{CASES}CASE-0002/").json()["status"] == "archived"


def test_dashboard_active_cases_leave_out_archived(owner, org):
    closed_case(org, 1, status="archived")
    Case.objects.create(organisation=org, number=2, title="Open", priority="low")

    summary = client_for(owner).get("/api/v1/dashboard/summary/").json()

    assert summary["active_cases"] == 1


def test_archiving_needs_a_member_of_the_organisation(outsider, org):
    closed_case(org)

    assert APIClient().post(f"{CASES}CASE-0001/status/", {"status": "archived"}, format="json").status_code == 401
    assert set_status(outsider, "CASE-0001", "archived").status_code == 404


def full_case(org, owner):
    case = Case.objects.create(organisation=org, number=1, title="Acme duplicates", priority="high", created_by=owner)
    make_alert(org, 2, case=case, explanation="Paid twice <b>bold</b>.")
    CaseEvent.objects.create(case=case, type=CaseEvent.CREATED, actor=owner, to_status="new", detail=case.title)
    CaseEvent.objects.create(case=case, type=CaseEvent.COMMENT_ADDED, actor=owner, detail="=1+1")
    CaseComment.objects.create(case=case, author=owner, text="First <i>look</i>")
    CaseComment.objects.create(case=case, author=owner, text="=HYPERLINK(\"http://evil\")")
    stored = StoredFile.objects.create(
        organisation=org, uploaded_by=owner, file_name="statement.pdf", content_type="application/pdf", size=1234,
        public_id="axiomvault/files/x",
    )
    CaseEvidence.objects.create(case=case, file=stored, added_by=owner)
    return case


def test_evidence_summary(owner, org):
    full_case(org, owner)

    response = client_for(owner).get(f"{CASES}case-0001/report/")

    assert response.status_code == 200
    body = response.json()
    assert set(body) == {"case", "alerts", "analysis", "related", "evidence", "comments", "history", "generated_at", "generated_by"}
    assert body["case"]["reference"] == "CASE-0001"
    assert [(a["reference"], a["explanation"]) for a in body["alerts"]] == [("ALT-0002", "Paid twice <b>bold</b>.")]
    assert len(body["alerts"][0]["flagged_records"]) == 2
    assert body["analysis"]["total_exposure"] == "84,700.00"
    assert body["related"]["vendors"][0]["id"] == "VND-8847"
    assert body["evidence"] == [{
        "file_name": "statement.pdf", "content_type": "application/pdf", "size": 1234, "source": "evidence",
        "added_by": {"id": str(owner.id), "full_name": owner.full_name, "profile_picture": None},
        "added_at": body["evidence"][0]["added_at"],
    }]
    assert [c["text"] for c in body["comments"]] == ["First <i>look</i>", "=HYPERLINK(\"http://evil\")"]
    assert [h["type"] for h in body["history"]] == ["created", "comment_added"]
    assert body["generated_by"]["id"] == str(owner.id)
    assert not AuditLog.objects.filter(action=AuditLog.CASE_REPORT_EXPORTED).exists()


def test_report_errors(outsider, org, owner):
    closed_case(org, status="archived")

    for path in ("report/", "report/pdf/", "report/xlsx/"):
        assert APIClient().get(f"{CASES}CASE-0001/{path}").status_code == 401
        assert client_for(outsider).get(f"{CASES}CASE-0001/{path}").status_code == 404
        assert client_for(owner).get(f"{CASES}CASE-0009/{path}").status_code == 404
        assert client_for(owner).get(f"{CASES}CASE-0001/{path}").status_code == 200  # Archived cases still open.
    assert not AuditLog.objects.filter(action=AuditLog.CASE_REPORT_EXPORTED, organisation=outsider.organisation).exists()


def test_pdf_export(owner, org):
    full_case(org, owner)

    response = client_for(owner).get(f"{CASES}CASE-0001/report/pdf/")

    assert response.status_code == 200
    assert response["Content-Type"] == "application/pdf"
    assert response["Content-Disposition"] == 'attachment; filename="CASE-0001-evidence-report.pdf"'
    assert response.content.startswith(b"%PDF")
    row = AuditLog.objects.get(action=AuditLog.CASE_REPORT_EXPORTED)
    assert (row.actor, row.details) == (owner, {"reference": "CASE-0001", "format": "pdf"})


def test_xlsx_export_keeps_text_as_text(owner, org):
    full_case(org, owner)

    response = client_for(owner).get(f"{CASES}CASE-0001/report/xlsx/")

    assert response.status_code == 200
    assert response["Content-Disposition"] == 'attachment; filename="CASE-0001-evidence-report.xlsx"'
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["Summary", "Alerts", "Flagged records", "Evidence", "Comments", "History"]
    assert book["Summary"]["B2"].value == "CASE-0001: Acme duplicates"
    assert book["Alerts"]["A2"].value == "ALT-0002"
    assert [row[0] for row in book["Flagged records"].iter_rows(min_row=2, values_only=True)] == ["ALT-0002", "ALT-0002"]
    assert book["Evidence"]["A2"].value == "statement.pdf"
    formula = book["Comments"]["C3"]
    assert (formula.value, formula.data_type) == ("=HYPERLINK(\"http://evil\")", "s")
    assert book["History"]["F3"].data_type == "s"
    assert AuditLog.objects.get(action=AuditLog.CASE_REPORT_EXPORTED).details == {"reference": "CASE-0001", "format": "xlsx"}


def test_pdf_escapes_markup_in_user_text():
    data = {
        "case": {"reference": "CASE-0001", "title": "<b>x</b> & <unclosed", "status": "new", "priority": "low",
                 "investigator": None, "reviewer": None, "created_at": "2026-10-15T09:00:00Z", "closed_at": None,
                 "description": "<para>"},
        "alerts": [], "evidence": [], "comments": [{"created_at": "2026-10-15T09:00:00Z", "author": None, "text": "<script>"}],
        "history": [], "related": {"employees": [], "vendors": [], "approvers": []},
        "analysis": {"total_exposure": "0.00", "duplicate_amount": "0.00", "risk_score": 0, "transaction_count": 0, "currency": "NGN", "source": "derived"},
        "generated_at": "2026-10-22T18:00:00Z", "generated_by": None,
    }

    assert ReportService.pdf(data).startswith(b"%PDF")
